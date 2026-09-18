"""
Sample DOCX Document Generator
Creates a sample company policy DOCX file for testing the Vectorless Groq Agent.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os


def create_sample_docx(output_path: str = "sample_company_policy.docx"):
    doc = docx.Document()

    # Title
    title = doc.add_heading("Acme Corp - Employee Handbook & Policies", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1
    doc.add_heading("1. Company Overview & Culture", level=1)
    doc.add_paragraph(
        "Acme Corp is a leading technology enterprise committed to innovation, integrity, and operational excellence. "
        "Founded in 2018, Acme Corp operates globally across North America, Europe, and Asia Pacific."
    )

    # Section 2
    doc.add_heading("2. Work Hours & Remote Work Policy", level=1)
    doc.add_paragraph(
        "Standard business hours are Monday through Friday, 9:00 AM to 5:00 PM local time. "
        "Core collaborative hours are set between 10:00 AM and 3:00 PM."
    )
    doc.add_heading("2.1 Hybrid Work Rules", level=2)
    p = doc.add_paragraph("Employees are eligible for hybrid work under the following guidelines:")
    doc.add_paragraph("• Employees must work in office at least 2 mandatory days per week (Tuesday & Thursday).", style='List Bullet')
    doc.add_paragraph("• Remote work from outside the country requires prior written approval from HR and Security.", style='List Bullet')
    doc.add_paragraph("• A annual home-office stipend of $500 is provided to full-time employees.", style='List Bullet')

    # Section 3
    doc.add_heading("3. Leave & Vacation Policy", level=1)
    doc.add_paragraph(
        "All full-time employees receive 20 days of paid vacation per calendar year, accrued monthly. "
        "In addition, employees receive 10 paid sick days and 12 official company holidays."
    )
    doc.add_heading("3.1 Parental Leave", level=2)
    doc.add_paragraph(
        "Acme Corp offers 16 weeks of fully paid parental leave for primary caregivers and 8 weeks for secondary caregivers, "
        "applicable following birth or adoption."
    )

    # Section 4
    doc.add_heading("4. Expense Allowances & Compensation Matrix", level=1)
    doc.add_paragraph(
        "Below is the standardized allowance and reimbursement cap matrix for travel and equipment:"
    )

    # Table
    table = doc.add_table(rows=5, cols=4)
    table.style = 'Table Grid'

    headers = ["Expense Category", "Tier 1 Cities Cap", "Tier 2 Cities Cap", "Approval Required"]
    hdr_cells = table.rows[0].cells
    for i, title_text in enumerate(headers):
        hdr_cells[i].text = title_text

    data = [
        ["Daily Meal Allowance", "$75 / day", "$50 / day", "Direct Manager"],
        ["Hotel Accommodation", "$250 / night", "$150 / night", "Direct Manager"],
        ["Client Entertainment", "$150 / event", "$100 / event", "Department Head"],
        ["Mobile & Data Stipend", "$80 / month", "$80 / month", "HR Automatic"],
    ]

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text

    # Section 5
    doc.add_heading("5. IT Security & Data Protection", level=1)
    doc.add_paragraph(
        "All laptops must have company-mandated VPN and antivirus software enabled at all times. "
        "Passwords must be changed every 90 days, with multi-factor authentication (MFA) strictly enforced on all accounts."
    )
    doc.add_paragraph(
        "For security incidents, contact security@acmecorp.com or call the emergency hotline at +1 (800) 555-0199."
    )

    doc.save(output_path)
    print(f"Sample DOCX successfully created at: {os.path.abspath(output_path)}")
    return output_path


if __name__ == "__main__":
    create_sample_docx()
